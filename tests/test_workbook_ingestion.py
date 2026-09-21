from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from openpyxl import Workbook

from engiverify.ingestion.workbook import analyze_workbook, read_sheet_cells

class WorkbookTests(unittest.TestCase):
    def make_workbook(self, root):
        path=Path(root)/"data.xlsx"
        wb=Workbook()
        ws=wb.active
        ws.title="Register"
        ws["A1"]="Tag"
        ws["B1"]="Qty"
        ws["A2"]="00123"
        ws["B2"]=4
        ws["C2"]="=B2*2"
        wb.save(path)
        wb.close()
        return path

    def test_structure_and_leading_zero_preservation(self):
        with TemporaryDirectory() as td:
            path=self.make_workbook(td)
            analysis=analyze_workbook(path)
            self.assertEqual(analysis["sheets"][0]["first_row_values"][:2],["Tag","Qty"])
            rows=read_sheet_cells(path,sheet_name="Register",min_row=2,max_row=2)
            self.assertEqual(rows[0]["cells"][0]["raw_value"],"00123")
            self.assertEqual(rows[0]["cells"][0]["data_type"],"s")
            self.assertEqual(rows[0]["cells"][2]["raw_value"],"=B2*2")

if __name__=="__main__":
    unittest.main()
