"""Workbook inspection that preserves source cell semantics."""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def analyze_workbook(path):
    source = Path(path)
    wb = load_workbook(source, read_only=True, data_only=False)
    try:
        sheets = []
        for ws in wb.worksheets:
            headers = []
            if ws.max_row >= 1:
                for cell in next(ws.iter_rows(min_row=1, max_row=1)):
                    headers.append(cell.value)
            sheets.append({
                "sheet_name": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "first_row_values": headers,
            })
        return {"sheet_count": len(sheets), "sheets": sheets}
    finally:
        wb.close()

def read_sheet_cells(path, *, sheet_name: str, min_row: int = 1, max_row: int | None = None):
    source = Path(path)
    formulas = load_workbook(source, read_only=True, data_only=False)
    cached = load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet_name not in formulas.sheetnames:
            raise KeyError(f"Worksheet not found: {sheet_name}")
        ws_formula = formulas[sheet_name]
        ws_cached = cached[sheet_name]
        end_row = max_row or ws_formula.max_row
        rows = []
        formula_rows = ws_formula.iter_rows(min_row=min_row, max_row=end_row)
        cached_rows = ws_cached.iter_rows(min_row=min_row, max_row=end_row)
        for formula_row, cached_row in zip(formula_rows, cached_rows):
            cells = []
            for a, b in zip(formula_row, cached_row):
                cells.append({
                    "coordinate": a.coordinate,
                    "column": get_column_letter(a.column),
                    "raw_value": a.value,
                    "cached_value": b.value,
                    "data_type": a.data_type,
                    "number_format": a.number_format,
                })
            rows.append({"row_number": formula_row[0].row if formula_row else None, "cells": cells})
        return rows
    finally:
        formulas.close()
        cached.close()
